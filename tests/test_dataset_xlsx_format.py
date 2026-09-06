from io import BytesIO

import openpyxl
import pytest

from agentgate.dataset.formats.xlsx import (
    HEADERS,
    SHEET_NAME,
    XlsxFormatError,
    dump,
    parse,
)
from agentgate.domain import (
    Case,
    CaseCategory,
    CaseDifficulty,
    CaseTurn,
    Equals,
    OutputExpectation,
    SkillRouteExpectation,
)


def case_payload() -> dict[str, object]:
    case = Case(
        id="loan-001",
        name="贷款申请",
        category=CaseCategory.BOUNDARY,
        difficulty=CaseDifficulty.HARD,
        tags=("loan", "中文"),
        notes="多轮场景",
        initial_state={"status": "new"},
        turns=(
            CaseTurn(
                id="turn-1",
                input={"message": "我想申请贷款"},
                expectations=(
                    SkillRouteExpectation(
                        id="route",
                        condition=Equals(expected="loan_approval"),
                    ),
                ),
                notes="识别意图",
            ),
            CaseTurn(
                id="turn-2",
                input={"message": "金额是20万元"},
                expectations=(
                    OutputExpectation(
                        id="output",
                        path="status",
                        condition=Equals(expected="submitted"),
                    ),
                ),
            ),
        ),
    )
    return case.model_dump(mode="json")


def workbook_bytes(
    rows: tuple[tuple[object, ...], ...],
    headers: tuple[str, ...] = HEADERS,
) -> bytes:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = SHEET_NAME
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    content = BytesIO()
    workbook.save(content)
    return content.getvalue()


def test_xlsx_runtime_uses_defused_xml() -> None:
    assert openpyxl.DEFUSEDXML


def test_xlsx_round_trip_preserves_current_multiturn_case() -> None:
    expected = case_payload()

    content = dump((expected,))
    parsed = parse(content)

    assert parsed == [expected]
    workbook = openpyxl.load_workbook(BytesIO(content), read_only=True)
    assert workbook.sheetnames == [SHEET_NAME]
    rows = list(workbook[SHEET_NAME].iter_rows(min_row=2, values_only=True))
    assert [row[0] for row in rows] == ["loan-001", "loan-001"]
    assert [row[8] for row in rows] == [1, 2]
    workbook.close()


def test_xlsx_accepts_simple_required_columns() -> None:
    content = workbook_bytes(
        (("case-1", "Simple", '{"message":"hello"}'),),
        ("case_id", "case_name", "input_json"),
    )

    parsed = parse(content)

    assert parsed == [
        {
            "id": "case-1",
            "name": "Simple",
            "category": "positive",
            "difficulty": "medium",
            "tags": [],
            "notes": "",
            "initial_state": {},
            "turns": [{"input": {"message": "hello"}, "expectations": [], "notes": ""}],
        }
    ]


def test_xlsx_groups_and_orders_turns_by_case_id() -> None:
    headers = ("case_id", "case_name", "turn_order", "input_json")
    content = workbook_bytes(
        (
            ("case-1", "Case", 2, '{"message":"second"}'),
            ("case-1", "Case", 1, '{"message":"first"}'),
        ),
        headers,
    )

    parsed = parse(content)

    assert [turn["input"]["message"] for turn in parsed[0]["turns"]] == [
        "first",
        "second",
    ]


def test_xlsx_reports_sheet_row_and_column() -> None:
    content = workbook_bytes(
        (("case-1", "Case", "not-json"),),
        ("case_id", "case_name", "input_json"),
    )

    with pytest.raises(XlsxFormatError) as error:
        parse(content)

    assert error.value.issues[0].sheet == SHEET_NAME
    assert error.value.issues[0].row == 2
    assert error.value.issues[0].column == "input_json"


@pytest.mark.parametrize(
    "input_json",
    ('{"message":"one","message":"two"}', '{"score":NaN}'),
)
def test_xlsx_rejects_lossy_json_values(input_json: str) -> None:
    content = workbook_bytes(
        (("case-1", "Case", input_json),),
        ("case_id", "case_name", "input_json"),
    )

    with pytest.raises(XlsxFormatError, match="invalid JSON"):
        parse(content)


def test_xlsx_rejects_formula_input() -> None:
    content = workbook_bytes(
        (("case-1", "Case", '=CONCAT("{", "}")'),),
        ("case_id", "case_name", "input_json"),
    )

    with pytest.raises(XlsxFormatError, match="formulas are not allowed"):
        parse(content)


def test_xlsx_export_keeps_formula_like_values_as_text() -> None:
    payload = case_payload()
    payload["id"] = "=1+1"
    payload["name"] = '=HYPERLINK("https://example.test")'

    content = dump((payload,))
    workbook = openpyxl.load_workbook(BytesIO(content), read_only=True, data_only=False)
    row = next(workbook[SHEET_NAME].iter_rows(min_row=2))

    assert row[0].data_type == "s"
    assert row[0].value == "=1+1"
    assert row[1].data_type == "s"
    workbook.close()


def test_xlsx_rejects_incomplete_turn_order_and_invalid_archive() -> None:
    content = workbook_bytes(
        (
            ("case-1", "Case", 1, '{"message":"first"}'),
            ("case-1", "Case", None, '{"message":"second"}'),
        ),
        ("case_id", "case_name", "turn_order", "input_json"),
    )

    with pytest.raises(XlsxFormatError, match="provided for every Turn"):
        parse(content)
    with pytest.raises(XlsxFormatError, match="not a valid XLSX archive"):
        parse(b"not an xlsx file")
