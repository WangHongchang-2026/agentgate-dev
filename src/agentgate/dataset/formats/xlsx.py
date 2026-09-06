"""One-sheet XLSX exchange for evaluation Case payloads."""

from __future__ import annotations

import json
from collections.abc import Mapping, Sequence
from dataclasses import dataclass
from io import BytesIO
from typing import Any
from zipfile import BadZipFile, ZipFile
from xml.etree.ElementTree import ParseError

from openpyxl import Workbook, load_workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.utils.exceptions import IllegalCharacterError, InvalidFileException


SHEET_NAME = "Cases"
HEADERS = (
    "case_id",
    "case_name",
    "category",
    "difficulty",
    "tags_json",
    "case_notes",
    "initial_state_json",
    "turn_id",
    "turn_order",
    "input_json",
    "expectations_json",
    "turn_notes",
)
REQUIRED_HEADERS = {"case_id", "case_name", "input_json"}
MAX_INPUT_BYTES = 10 * 1024 * 1024
MAX_ROWS = 10_000
MAX_ARCHIVE_ENTRIES = 2_000
MAX_UNCOMPRESSED_BYTES = 100 * 1024 * 1024
MAX_ENTRY_BYTES = 50 * 1024 * 1024
MAX_COMPRESSION_RATIO = 200
MAX_ISSUES = 200


@dataclass(frozen=True, slots=True)
class XlsxIssue:
    """One workbook problem located for user correction."""

    sheet: str
    row: int | None
    column: str | None
    message: str


class XlsxFormatError(ValueError):
    """Bounded collection of XLSX syntax and structure problems."""

    def __init__(self, issues: Sequence[XlsxIssue]) -> None:
        self.issues = tuple(issues[:MAX_ISSUES])
        details = "; ".join(
            f"{issue.sheet}"
            f"{f' row {issue.row}' if issue.row is not None else ''}"
            f"{f' column {issue.column}' if issue.column else ''}: {issue.message}"
            for issue in self.issues
        )
        super().__init__(details or "XLSX validation failed")


def _issue(
    issues: list[XlsxIssue],
    row: int | None,
    column: str | None,
    message: str,
) -> None:
    if len(issues) < MAX_ISSUES:
        issues.append(XlsxIssue(SHEET_NAME, row, column, message))


def _validate_archive(source: bytes) -> None:
    if len(source) > MAX_INPUT_BYTES:
        raise XlsxFormatError(
            (XlsxIssue(SHEET_NAME, None, None, "file exceeds the 10 MiB limit"),)
        )
    try:
        with ZipFile(BytesIO(source)) as archive:
            entries = archive.infolist()
    except (BadZipFile, OSError) as exc:
        raise XlsxFormatError(
            (XlsxIssue(SHEET_NAME, None, None, "file is not a valid XLSX archive"),)
        ) from exc

    issues: list[XlsxIssue] = []
    names = {entry.filename for entry in entries}
    if "[Content_Types].xml" not in names:
        _issue(issues, None, None, "XLSX content types are missing")
    if len(entries) > MAX_ARCHIVE_ENTRIES:
        _issue(issues, None, None, "archive contains too many entries")
    if sum(entry.file_size for entry in entries) > MAX_UNCOMPRESSED_BYTES:
        _issue(issues, None, None, "archive expands beyond the 100 MiB limit")

    forbidden = {
        "macros": any(name.endswith("vbaProject.bin") for name in names),
        "external links": any(name.startswith("xl/externalLinks/") for name in names),
        "embedded objects": any(name.startswith("xl/embeddings/") for name in names),
        "ActiveX controls": any(name.startswith("xl/activeX/") for name in names),
        "data connections": "xl/connections.xml" in names,
    }
    for label, present in forbidden.items():
        if present:
            _issue(issues, None, None, f"workbook {label} are not allowed")

    for entry in entries:
        if entry.flag_bits & 0x1:
            _issue(issues, None, None, "encrypted archive entries are not allowed")
        if entry.file_size > MAX_ENTRY_BYTES:
            _issue(issues, None, None, f"archive entry is too large: {entry.filename}")
        if entry.file_size / max(entry.compress_size, 1) > MAX_COMPRESSION_RATIO:
            _issue(
                issues,
                None,
                None,
                f"archive entry compression ratio is too high: {entry.filename}",
            )
    if issues:
        raise XlsxFormatError(issues)


def _parse_json_cell(
    value: Any,
    default: Any,
    expected_type: type,
    row: int,
    column: str,
    issues: list[XlsxIssue],
    *,
    required: bool = False,
) -> Any:
    if value is None or value == "":
        if required:
            _issue(issues, row, column, "value is required")
        return default
    if not isinstance(value, str):
        _issue(issues, row, column, "must contain JSON text")
        return default
    try:
        parsed = json.loads(
            value,
            object_pairs_hook=_object_without_duplicates,
            parse_constant=_reject_non_finite_number,
        )
    except (json.JSONDecodeError, ValueError) as exc:
        message = exc.msg if isinstance(exc, json.JSONDecodeError) else str(exc)
        _issue(issues, row, column, f"invalid JSON: {message}")
        return default
    if not isinstance(parsed, expected_type):
        _issue(issues, row, column, f"must contain a JSON {expected_type.__name__}")
        return default
    return parsed


def _object_without_duplicates(pairs: list[tuple[str, Any]]) -> dict[str, Any]:
    result: dict[str, Any] = {}
    for key, value in pairs:
        if key in result:
            raise ValueError(f"duplicate object key: {key}")
        result[key] = value
    return result


def _reject_non_finite_number(value: str) -> None:
    raise ValueError(f"non-finite number is not allowed: {value}")


def _parse_text(
    value: Any,
    row: int,
    column: str,
    issues: list[XlsxIssue],
    *,
    required: bool = False,
    default: str = "",
) -> str:
    if value is None or (isinstance(value, str) and not value.strip()):
        if required:
            _issue(issues, row, column, "value is required")
        return default
    if not isinstance(value, str):
        _issue(issues, row, column, "must be text")
        return default
    return value.strip() if required else value


def _parse_order(value: Any, row: int, issues: list[XlsxIssue]) -> int | None:
    if value is None or value == "":
        return None
    if isinstance(value, bool):
        _issue(issues, row, "turn_order", "must be a positive integer")
        return None
    if isinstance(value, int):
        order = value
    elif isinstance(value, float) and value.is_integer():
        order = int(value)
    else:
        try:
            order = int(value)
        except (TypeError, ValueError):
            _issue(issues, row, "turn_order", "must be a positive integer")
            return None
    if order < 1:
        _issue(issues, row, "turn_order", "must be a positive integer")
        return None
    return order


def parse(source: bytes) -> list[dict[str, object]]:
    """Parse one Cases worksheet into ordered plain Case payloads."""

    _validate_archive(source)
    try:
        workbook = load_workbook(
            BytesIO(source), read_only=True, data_only=False, keep_links=False
        )
    except (InvalidFileException, KeyError, OSError, ValueError) as exc:
        raise XlsxFormatError(
            (XlsxIssue(SHEET_NAME, None, None, "workbook cannot be opened"),)
        ) from exc

    try:
        if SHEET_NAME not in workbook.sheetnames:
            raise XlsxFormatError(
                (XlsxIssue(SHEET_NAME, None, None, "required worksheet is missing"),)
            )
        sheet = workbook[SHEET_NAME]
        rows = iter(sheet.iter_rows())
        header_cells = tuple(next(rows, ()))
        header = tuple(cell.value for cell in header_cells)
        issues: list[XlsxIssue] = []

        if any(cell.data_type == "f" for cell in header_cells):
            _issue(issues, 1, None, "formulas are not allowed")
        if any(not isinstance(value, str) for value in header):
            _issue(issues, 1, None, "headers must be text")
        names = tuple(value for value in header if isinstance(value, str))
        if len(names) != len(set(names)):
            _issue(issues, 1, None, "headers must not contain duplicates")
        for name in names:
            if name not in HEADERS:
                _issue(issues, 1, name, "unexpected header")
        for required in sorted(REQUIRED_HEADERS - set(names)):
            _issue(issues, 1, required, "required header is missing")
        if issues:
            raise XlsxFormatError(issues)

        grouped: dict[str, list[tuple[int, dict[str, Any]]]] = {}
        for row_number, cells in enumerate(rows, start=2):
            if row_number > MAX_ROWS + 1:
                _issue(issues, row_number, None, "worksheet exceeds 10,000 data rows")
                break
            values: dict[str, Any] = {}
            for index, name in enumerate(names):
                cell = cells[index] if index < len(cells) else None
                if cell is not None and cell.data_type == "f":
                    _issue(issues, row_number, name, "formulas are not allowed")
                    values[name] = None
                else:
                    values[name] = cell.value if cell is not None else None
            if all(value is None or value == "" for value in values.values()):
                continue

            case_id = _parse_text(
                values.get("case_id"), row_number, "case_id", issues, required=True
            )
            row = {
                "case_id": case_id,
                "case_name": _parse_text(
                    values.get("case_name"),
                    row_number,
                    "case_name",
                    issues,
                    required=True,
                ),
                "category": _parse_text(
                    values.get("category"),
                    row_number,
                    "category",
                    issues,
                    default="positive",
                ),
                "difficulty": _parse_text(
                    values.get("difficulty"),
                    row_number,
                    "difficulty",
                    issues,
                    default="medium",
                ),
                "tags": _parse_json_cell(
                    values.get("tags_json"), [], list, row_number, "tags_json", issues
                ),
                "notes": _parse_text(
                    values.get("case_notes"), row_number, "case_notes", issues
                ),
                "initial_state": _parse_json_cell(
                    values.get("initial_state_json"),
                    {},
                    dict,
                    row_number,
                    "initial_state_json",
                    issues,
                ),
                "turn_id": _parse_text(
                    values.get("turn_id"), row_number, "turn_id", issues
                ),
                "turn_order": _parse_order(
                    values.get("turn_order"), row_number, issues
                ),
                "input": _parse_json_cell(
                    values.get("input_json"),
                    {},
                    dict,
                    row_number,
                    "input_json",
                    issues,
                    required=True,
                ),
                "expectations": _parse_json_cell(
                    values.get("expectations_json"),
                    [],
                    list,
                    row_number,
                    "expectations_json",
                    issues,
                ),
                "turn_notes": _parse_text(
                    values.get("turn_notes"), row_number, "turn_notes", issues
                ),
            }
            if case_id:
                grouped.setdefault(case_id, []).append((row_number, row))

        output: list[dict[str, object]] = []
        case_fields = (
            "case_name",
            "category",
            "difficulty",
            "tags",
            "notes",
            "initial_state",
        )
        case_field_columns = {
            "case_name": "case_name",
            "category": "category",
            "difficulty": "difficulty",
            "tags": "tags_json",
            "notes": "case_notes",
            "initial_state": "initial_state_json",
        }
        for case_id, case_rows in grouped.items():
            first_number, first = case_rows[0]
            for row_number, row in case_rows[1:]:
                for field in case_fields:
                    if row[field] != first[field]:
                        _issue(
                            issues,
                            row_number,
                            case_field_columns[field],
                            f"conflicts with row {first_number} for Case {case_id}",
                        )

            orders = [row["turn_order"] for _, row in case_rows]
            if all(order is None for order in orders):
                ordered_rows = case_rows
            elif any(order is None for order in orders):
                _issue(
                    issues,
                    first_number,
                    "turn_order",
                    "must be provided for every Turn in this Case",
                )
                ordered_rows = case_rows
            else:
                expected = list(range(1, len(case_rows) + 1))
                actual = sorted(int(order) for order in orders if order is not None)
                if actual != expected:
                    _issue(
                        issues,
                        first_number,
                        "turn_order",
                        "must be unique and contiguous from 1",
                    )
                ordered_rows = sorted(
                    case_rows, key=lambda item: int(item[1]["turn_order"] or 0)
                )

            turns: list[dict[str, object]] = []
            turn_ids: set[str] = set()
            for row_number, row in ordered_rows:
                turn: dict[str, object] = {
                    "input": row["input"],
                    "expectations": row["expectations"],
                    "notes": row["turn_notes"],
                }
                turn_id = str(row["turn_id"])
                if turn_id:
                    if turn_id in turn_ids:
                        _issue(issues, row_number, "turn_id", "duplicate within Case")
                    turn_ids.add(turn_id)
                    turn["id"] = turn_id
                turns.append(turn)
            output.append(
                {
                    "id": case_id,
                    "name": first["case_name"],
                    "category": first["category"],
                    "difficulty": first["difficulty"],
                    "tags": first["tags"],
                    "notes": first["notes"],
                    "initial_state": first["initial_state"],
                    "turns": turns,
                }
            )

        if issues:
            raise XlsxFormatError(issues)
        return output
    except XlsxFormatError:
        raise
    except (BadZipFile, EOFError, KeyError, OSError, ParseError, ValueError) as exc:
        raise XlsxFormatError(
            (XlsxIssue(SHEET_NAME, None, None, "workbook content is invalid"),)
        ) from exc
    finally:
        workbook.close()


def _json_text(value: Any) -> str:
    return json.dumps(
        value, ensure_ascii=False, sort_keys=True, separators=(",", ":"), allow_nan=False
    )


def _text_cell(sheet: Any, value: Any) -> Any:
    if not isinstance(value, str):
        return value
    cell = WriteOnlyCell(sheet, value=value)
    cell.data_type = "s"
    return cell


def dump(cases: Sequence[Mapping[str, Any]]) -> bytes:
    """Write plain Case payloads to the one-sheet XLSX format."""

    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet(SHEET_NAME)
    sheet.freeze_panes = "A2"
    sheet.append(tuple(_text_cell(sheet, value) for value in HEADERS))
    try:
        for case in cases:
            turns = case.get("turns")
            if not isinstance(turns, (list, tuple)) or not turns:
                raise ValueError("XLSX export Case requires at least one Turn")
            for order, turn in enumerate(turns, start=1):
                if not isinstance(turn, Mapping):
                    raise ValueError("XLSX export Turn must be an object")
                row = (
                    case.get("id", ""),
                    case.get("name", ""),
                    case.get("category", "positive"),
                    case.get("difficulty", "medium"),
                    _json_text(case.get("tags", [])),
                    case.get("notes", ""),
                    _json_text(case.get("initial_state", {})),
                    turn.get("id", ""),
                    order,
                    _json_text(turn.get("input", {})),
                    _json_text(turn.get("expectations", [])),
                    turn.get("notes", ""),
                )
                sheet.append(tuple(_text_cell(sheet, value) for value in row))
        content = BytesIO()
        workbook.save(content)
    except (IllegalCharacterError, TypeError, ValueError) as exc:
        raise XlsxFormatError(
            (XlsxIssue(SHEET_NAME, None, None, f"cannot export workbook: {exc}"),)
        ) from exc
    finally:
        workbook.close()
    return content.getvalue()
